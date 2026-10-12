from __future__ import annotations

import html
import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "outputs/translation-basis-20260603/smicut-cutting-data-translation-basis.xlsx"
FILES = [
    ROOT / "Translation/Spanska/smicut cutting data app-es.html",
    ROOT / "Translation/Spanska/smicut material categories-es.html",
]
EXPECTED_API_LANGUAGE = "es"


def load_mapping() -> dict[str, str]:
    wb = load_workbook(WORKBOOK, data_only=False)
    ws = wb["Translations"]
    mapping: dict[str, str] = {}
    for row in range(2, ws.max_row + 1):
        english = str(ws.cell(row, 1).value or "")
        spanish = str(ws.cell(row, 3).value or "")
        if english and spanish and english != spanish:
            mapping[english] = spanish
            mapping[re.sub(r"(?<=\d)\.(?=\d)", ",", english)] = spanish
    return mapping


def js_unescape(value: str) -> str:
    return (
        value.replace("\\'", "'")
        .replace('\\"', '"')
        .replace("\\n", "\n")
        .replace("\\r", "\r")
        .replace("\\\\", "\\")
    )


def validate_file(path: Path, mapping: dict[str, str]) -> dict[str, object]:
    content = path.read_text(encoding="utf-8")
    non_script = re.sub(r"<script\b[^>]*>.*?</script>", "", content, flags=re.I | re.S)
    leftovers: list[dict[str, str]] = []

    for match in re.finditer(r">([^<>]+)<", non_script):
        text = html.unescape(match.group(1).strip())
        if text in mapping:
            leftovers.append({"kind": "text", "value": text, "expected": mapping[text]})

    for match in re.finditer(r"\b(?:aria-label|placeholder|title|alt)=(\"|')(.*?)(\1)", non_script):
        text = html.unescape(match.group(2))
        if text in mapping:
            leftovers.append({"kind": "attribute", "value": text, "expected": mapping[text]})

    for script_match in re.finditer(r"<script\b[^>]*>(.*?)</script>", content, flags=re.I | re.S):
        script = script_match.group(1)
        for match in re.finditer(r"(['\"])((?:\\.|(?!\1).)*?)(\1)", script):
            text = js_unescape(match.group(2))
            if text in mapping:
                leftovers.append({"kind": "js", "value": text, "expected": mapping[text]})

    en_links = re.findall(r"https://smicut\.com/en/[^\"'\s<]*", content)
    api_language_match = re.search(
        r'\blanguage:\s*["\']([A-Za-z0-9_-]+)["\']\s*,\s*\n\s*history:\s*history',
        content,
    )
    api_language = api_language_match.group(1) if api_language_match else None
    return {
        "file": str(path.relative_to(ROOT)),
        "leftover_translatable_strings": leftovers,
        "english_smicut_links": en_links,
        "material_selector_api_language": api_language,
        "expected_material_selector_api_language": EXPECTED_API_LANGUAGE if api_language else None,
    }


def main() -> None:
    mapping = load_mapping()
    result = [validate_file(path, mapping) for path in FILES]
    print(json.dumps(result, ensure_ascii=False, indent=2))
    failed = any(
        item["leftover_translatable_strings"]
        or item["english_smicut_links"]
        or (
            item["material_selector_api_language"]
            and item["material_selector_api_language"] != EXPECTED_API_LANGUAGE
        )
        for item in result
    )
    raise SystemExit(1 if failed else 0)


if __name__ == "__main__":
    main()
