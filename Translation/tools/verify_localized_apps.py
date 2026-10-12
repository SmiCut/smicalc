from __future__ import annotations

import argparse
import html
import importlib.util
import json
import re
import subprocess
import os
import shutil
import tempfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
GENERATOR = ROOT / "Translation/tools/generate_localized_html.py"
WORKBOOK = ROOT / "outputs/translation-basis-20260603/smicut-cutting-data-translation-basis.xlsx"
NODE = Path(os.environ.get("NODE_BINARY") or shutil.which("node") or "node")
ENGLISH_APP = ROOT / "Translation/English/smicut cutting data app-eng.html"
MAIN_APP = ROOT / "smicut cutting data app.html"
MATERIAL_CATEGORY_KEYS_FILE = ROOT / "Translation/material-category-keys.json"


def load_generator():
    spec = importlib.util.spec_from_file_location("smicalc_localization_generator", GENERATOR)
    if not spec or not spec.loader:
        raise RuntimeError(f"Could not load generator from {GENERATOR}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


FALSE_FRIEND_PATTERNS = {
    "thread/topic/wire": [
        r"\bThread-Profil\b",
        r"\bProfil de fil\b",
        r"\bSzálprofil\b",
        r"\bDiscussieprofiel\b",
        r"\bDraadprofiel\b",
        r"\bProfil wątku\b",
        r"Профиль темы",
        r"螺纹轮廓",
        r"螺紋輪廓",
        r"スレッド\s*プロファイル",
        r"스레드 프로필",
        r"\bTrådprofil\b",
    ],
    "flute-as-instrument": [
        r"flûtes?",
        r"cannelures?",
        r"flauti?",
        r"fuvolák?",
        r"\bFlöten?\b",
        r"\bfluiten?\b",
        r"seruling",
        r"笛",
        r"フルート",
    ],
    "pass-vs-thread": [
        r"menetek szám[aá]t?",
        r"Nombre de passages",
        r"\bpassages\b",
        r"przejścia gwintu",
    ],
    "insert-action": [
        r"\bEinfügen\b",
        r"\bInsérer\b",
        r"\bInvoegen\b",
        r"\bWstaw\b",
        r"Вставить",
        r"插入",
        r"挿入",
        r"삽입",
    ],
    "cutting-off-data": [
        r"данные резки",
        r"для резки",
        r"切断データ",
        r"切断条件",
        r"절단 데이터",
        r"절단 조건",
        r"切削資料",
        r"Parameter Cutting",
    ],
    "grade-rank-quality": [
        r"EmpFehlung",
        r"等级",
        r"等級",
        r"등급",
        r"\bfokozat\b",
    ],
}


def app_path_for(locale: str, config: dict[str, str]) -> Path:
    return ROOT / "Translation" / config["folder"] / f"smicut cutting data app-{config['suffix']}.html"


def selected_locales(all_locales: dict[str, dict[str, str]], raw: str) -> dict[str, dict[str, str]]:
    if raw == "all":
        return all_locales
    wanted = {item.strip() for item in raw.split(",") if item.strip()}
    unknown = wanted - set(all_locales)
    if unknown:
        raise SystemExit(f"Unknown locale(s): {', '.join(sorted(unknown))}")
    return {locale: all_locales[locale] for locale in sorted(wanted)}


def check_api_language(locales: dict[str, dict[str, str]]) -> list[str]:
    issues: list[str] = []
    for locale, config in locales.items():
        path = app_path_for(locale, config)
        if not path.exists():
            issues.append(f"{locale}: missing app file {path.relative_to(ROOT)}")
            continue
        text = path.read_text(encoding="utf-8")
        expected = config["api_language"]
        if f'language: "{expected}"' not in text and f"language: '{expected}'" not in text:
            found = re.findall(r"language:\s*['\"]([^'\"]+)['\"]", text)
            issues.append(f"{locale}: expected material selector language {expected}, found {found}")
    return issues


def extract_material_keys(text: str) -> list[str]:
    match = re.search(
        r"const\s+materialData\s*=\s*\[(?P<body>.*?)\];",
        text,
        flags=re.S,
    )
    if not match:
        return []
    return re.findall(r'\bkey\s*:\s*["\']([^"\']+)["\']', match.group("body"))


def check_material_keys(locales: dict[str, dict[str, str]]) -> list[str]:
    issues: list[str] = []
    expected = json.loads(MATERIAL_CATEGORY_KEYS_FILE.read_text(encoding="utf-8"))
    if len(expected) != 30 or len(set(expected)) != 30:
        return [f"Canonical key file must contain 30 unique material keys, found {len(expected)}"]

    for label, path in (("English source", ENGLISH_APP), ("main app", MAIN_APP)):
        actual = extract_material_keys(path.read_text(encoding="utf-8"))
        if actual != expected:
            issues.append(f"{label}: material keys differ from the canonical key file")

    for locale, config in locales.items():
        path = app_path_for(locale, config)
        if not path.exists():
            continue
        actual = extract_material_keys(path.read_text(encoding="utf-8"))
        if actual != expected:
            issues.append(f"{locale}: material keys differ from the English source")
    return issues


def check_material_selector_history(locales: dict[str, dict[str, str]]) -> list[str]:
    issues: list[str] = []
    for locale, config in locales.items():
        path = app_path_for(locale, config)
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        match = re.search(
            r"function\s+getRequestHistory\s*\(\)\s*\{(?P<body>.*?)\n\s*\}\s*\n\s*async\s+function\s+sendMessage",
            text,
            flags=re.S,
        )
        if not match:
            issues.append(f"{locale}: missing getRequestHistory")
            continue
        body = match.group("body")
        if 'message.role === "assistant"' not in body or ".slice(-12)" not in body:
            issues.append(f"{locale}: material selector history does not retain assistant replies")
    return issues


def check_translation_workbook_canonical_labels(
    generator,
    locales: dict[str, dict[str, str]],
) -> list[str]:
    issues: list[str] = []
    wb = load_workbook(WORKBOOK, read_only=True, data_only=False)
    ws = wb["Translations"]
    headers = [cell.value for cell in ws[1]]
    header_indexes = {
        header: index
        for index, header in enumerate(headers)
        if isinstance(header, str)
    }
    rows_by_english: dict[str, list[tuple[int, tuple]]] = {}
    for row_number, row in enumerate(
        ws.iter_rows(min_row=2, max_col=len(headers), values_only=True),
        start=2,
    ):
        source = row[0]
        if isinstance(source, str):
            rows_by_english.setdefault(source, []).append((row_number, row))

    def compare(source: str, locale: str, expected: str) -> None:
        matches = rows_by_english.get(source, [])
        if not matches:
            issues.append(
                f"translation workbook: missing English source row {source!r}"
            )
            return

        target_column = locales[locale]["target_column"]
        column_index = header_indexes.get(target_column)
        if column_index is None:
            issues.append(
                f"translation workbook: missing target column {target_column!r}"
            )
            return

        for row_number, row in matches:
            actual = row[column_index]
            if actual != expected:
                issues.append(
                    f"translation workbook row {row_number}, {target_column}: "
                    f"expected {expected!r}, found {actual!r}"
                )

    for locale in locales:
        for key, source in generator.THREAD_PROFILE_LABELS["en"].items():
            compare(source, locale, generator.THREAD_PROFILE_LABELS[locale][key])
        for key, source in generator.TAP_TYPE_LABELS["en"].items():
            compare(source, locale, generator.TAP_TYPE_LABELS[locale][key])

    for locale, overrides in generator.UI_TRANSLATION_OVERRIDES.items():
        if locale in locales:
            for source, expected in overrides.items():
                compare(source, locale, expected)

    if "ja" in locales:
        compare("➜ SmiProg", "ja", "➜ スミプログ")
        compare(
            "SmiProg - Software for Thread Milling",
            "ja",
            "スミプログ - ねじフライス加工用ソフトウェア",
        )
        japanese_column = header_indexes.get(locales["ja"]["target_column"])
        if japanese_column is not None:
            japanese_brand_names = {
                "SmiProg": "スミプログ",
                "SmiCalc": "スミカルク",
                "SmiCut": "スミカット",
            }
            for row_number, row in (
                item
                for items in rows_by_english.values()
                for item in items
            ):
                value = row[japanese_column]
                if not isinstance(value, str):
                    continue
                check_value = value.replace("SmiCut AB", "")
                for latin_name, japanese_name in japanese_brand_names.items():
                    if latin_name in check_value:
                        issues.append(
                            f"translation workbook row {row_number}, "
                            f"{locales['ja']['target_column']}: "
                            f"use {japanese_name} instead of {latin_name} in Japanese user-facing text"
                        )

    return issues


def check_japanese_brand_names(
    locales: dict[str, dict[str, str]],
) -> list[str]:
    if "ja" not in locales:
        return []

    path = app_path_for("ja", locales["ja"])
    if not path.exists():
        return []

    text = path.read_text(encoding="utf-8")
    user_facing = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    # Standalone JavaScript comment lines are developer notes, not UI labels.
    user_facing = re.sub(r"(?m)^\s*//[^\n]*(?:\n|$)", "", user_facing)
    user_facing = user_facing.replace("SmiCut AB", "")
    issues: list[str] = []
    for latin_name, japanese_name in {
        "SmiProg": "スミプログ",
        "SmiCalc": "スミカルク",
        "SmiCut": "スミカット",
    }.items():
        if latin_name in user_facing:
            issues.append(
                f"ja: use {japanese_name} instead of {latin_name} "
                "in Japanese user-facing app text"
            )
    return issues


def extract_thread_profile_labels(text: str) -> dict[str, str]:
    match = re.search(r"threadProfiles:\s*\{(?P<body>.*?)\n\s*\},", text, flags=re.S)
    if not match:
        return {}
    return dict(re.findall(r"(\w+):\s*'([^']*)'", match.group("body")))


def check_thread_profile_labels(generator, locales: dict[str, dict[str, str]]) -> list[str]:
    issues: list[str] = []
    npt_nationality_terms = {
        "sv": r"amerikansk",
        "es": r"americana",
        "es-gl": r"americana",
        "id": r"Amerika",
        "da": r"amerikansk",
        "de": r"\bUS-",
        "fr-be": r"américain",
        "fr": r"américain",
        "it": r"americana",
        "hu": r"amerikai",
        "nl-be": r"Amerikaanse",
        "nl": r"Amerikaanse",
        "pl": r"amerykański",
        "ru": r"американская",
        "zh": r"美制",
        "zh-tw": r"美制",
        "ja": r"ナショナル",
    }
    targets = [
        ("English source", ENGLISH_APP, generator.THREAD_PROFILE_LABELS["en"]),
        ("main app", MAIN_APP, generator.THREAD_PROFILE_LABELS["en"]),
    ]
    targets.extend(
        (
            locale,
            app_path_for(locale, config),
            generator.THREAD_PROFILE_LABELS[locale],
        )
        for locale, config in locales.items()
    )

    for label, path, expected in targets:
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        actual = extract_thread_profile_labels(text)
        if actual != expected:
            issues.append(f"{label}: JavaScript thread-profile labels differ from the canonical labels")

        if "DIN 40430" not in expected["pg"]:
            issues.append(
                f"{label}: PG profile label must include the standard DIN 40430"
            )

        copy_prompt = re.search(
            r"completeInputsBeforeCopying:\s*'(?P<value>[^']*)'",
            text,
        )
        if copy_prompt and copy_prompt.group("value") in expected.values():
            issues.append(
                f"{label}: copy-validation prompt was overwritten by a thread-profile label"
            )

        forbidden_npt_term = npt_nationality_terms.get(label)
        if forbidden_npt_term and any(
            re.search(
                forbidden_npt_term,
                expected[key],
                flags=re.IGNORECASE,
            )
            for key in ("npt", "nptf")
        ):
            issues.append(
                f"{label}: NPT/NPTF label still contains a nationality term"
            )

        if label == "sv":
            ordinary_profiles = (
                "metricCoarse",
                "metricFine",
                "unc",
                "unf",
                "unef",
                "unConstant",
            )
            if any("gänga" in expected[key].casefold() for key in ordinary_profiles):
                issues.append(
                    "sv: ordinary metric/Unified profile labels must use compact forms without 'gänga'"
                )

            pipe_profiles = ("g", "gOther", "bspt", "npt", "nptf")
            if any("rörgänga" not in expected[key].casefold() for key in pipe_profiles):
                issues.append(
                    "sv: G/BSPT/NPT/NPTF profile labels must retain the full word 'rörgänga'"
                )
            if "pansarrörsgänga" not in expected["pg"].casefold():
                issues.append(
                    "sv: PG profile label must retain the full word 'pansarrörsgänga'"
                )

        compact_pipe_terms = {
            "fr-be": "filetage",
            "fr": "filetage",
            "it": "filettatura",
        }
        if label in compact_pipe_terms:
            forbidden_term = compact_pipe_terms[label]
            standard_pipe_profiles = ("g", "gOther", "bspt", "npt", "nptf")
            if any(
                forbidden_term in expected[key].casefold()
                for key in standard_pipe_profiles
            ):
                issues.append(
                    f"{label}: standard pipe-profile labels must omit the contextually redundant word {forbidden_term!r}"
                )
            if label == "it" and "filettatura" not in expected["pg"].casefold():
                issues.append(
                    f"{label}: PG profile label must retain the full technical word 'filettatura'"
                )

        if label in {"nl-be", "nl"}:
            standard_pipe_profiles = ("g", "gOther", "bspt", "npt", "nptf")
            if any(
                "pijpschroefdraad" in expected[key].casefold()
                for key in standard_pipe_profiles
            ):
                issues.append(
                    f"{label}: standard pipe-profile labels must use the compact industry term 'pijpdraad'"
                )
            if any(
                "pijpdraad" not in expected[key].casefold()
                for key in standard_pipe_profiles
            ):
                issues.append(
                    f"{label}: standard pipe-profile labels are missing the term 'pijpdraad'"
                )

        lowercase_profile_starts = {
            "pl": (
                "metricCoarse",
                "metricFine",
                "unc",
                "unf",
                "unef",
                "unConstant",
                "g",
                "gOther",
                "bspt",
                "pg",
                "npt",
            ),
            "ru": tuple(expected),
        }
        if label in lowercase_profile_starts:
            for key in lowercase_profile_starts[label]:
                description = expected[key].split(" - ", 1)[-1]
                first_letter = next(
                    (character for character in description if character.isalpha()),
                    "",
                )
                if first_letter and not first_letter.islower():
                    issues.append(
                        f"{label}: {key} description must start with a lowercase letter"
                    )

        if label == "ko" and "나사 (" in expected["gOther"]:
            issues.append(
                "ko: manual-note parentheses must not have a preceding space"
            )

        for key, expected_label in expected.items():
            visible = [
                html.unescape(value)
                for value in re.findall(
                    rf'<option value="{re.escape(key)}">([^<]*)</option>',
                    text,
                )
            ]
            if len(visible) != 2 or any(value != expected_label for value in visible):
                issues.append(
                    f"{label}: expected two visible {key!r} options with label {expected_label!r}, found {visible}"
                )
    return issues


def check_material_label_deduplication(locales: dict[str, dict[str, str]]) -> list[str]:
    issues: list[str] = []
    targets = [
        ("English source", ENGLISH_APP),
        ("main app", MAIN_APP),
    ]
    targets.extend(
        (locale, app_path_for(locale, config))
        for locale, config in locales.items()
    )
    marker = "group.toLocaleLowerCase() === descriptor.toLocaleLowerCase()"

    for label, path in targets:
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        if marker not in text:
            issues.append(
                f"{label}: material labels do not suppress identical group and descriptor text"
            )
    return issues


def check_adjustable_input_accessibility(locales: dict[str, dict[str, str]]) -> list[str]:
    issues: list[str] = []
    targets = [
        ("English source", ENGLISH_APP),
        ("main app", MAIN_APP),
    ]
    targets.extend(
        (locale, app_path_for(locale, config))
        for locale, config in locales.items()
    )

    for label, path in targets:
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        inputs = re.findall(
            r'<input\b(?=[^>]*\bclass="[^"]*\bscd-adjustable-input\b[^"]*")[^>]*>',
            text,
        )
        if len(inputs) != 9:
            issues.append(
                f"{label}: expected 9 adjustable inputs, found {len(inputs)}"
            )
            continue

        for input_html in inputs:
            input_id = re.search(r'\bid="([^"]+)"', input_html)
            input_name = input_id.group(1) if input_id else "unknown"
            placeholder = re.search(r'\bplaceholder="([^"]*)"', input_html)
            if placeholder and placeholder.group(1):
                issues.append(
                    f"{label}: {input_name} has visible placeholder text"
                )
            for attr in ("aria-label", "title"):
                match = re.search(rf'\b{re.escape(attr)}="([^"]*)"', input_html)
                if not match or not match.group(1).strip():
                    issues.append(
                        f"{label}: {input_name} lacks a non-empty {attr}"
                    )
    return issues


def check_cjk_typography(generator, locales: dict[str, dict[str, str]]) -> list[str]:
    issues: list[str] = []
    required_tokens = (
        'id="scd-cjk-typography"',
        'class="scd-cjk-ui"',
        "#smicut-cutting-data-app.scd-cjk-ui .scd-control",
        "#smicut-cutting-data-app.scd-cjk-ui .scd-select option",
        "#smicut-cutting-data-app.scd-cjk-ui .scd-number",
        "#smicut-cutting-data-app.scd-cjk-ui .scd-choice span",
        "#smicut-cutting-data-app.scd-cjk-ui .scd-input-summary .scd-result-value",
        "#smicut-cutting-data-app.scd-cjk-ui #smicut-chat-widget .sms-message",
        "font-weight: 500 !important;",
        "font-weight: 600 !important;",
        "@media (min-width: 760px)",
    )

    for locale, config in locales.items():
        path = app_path_for(locale, config)
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        if locale in generator.CJK_UI_LOCALES:
            for token in required_tokens:
                if token not in text:
                    issues.append(
                        f"{locale}: missing CJK typography token {token!r}"
                    )
        elif 'id="scd-cjk-typography"' in text or 'class="scd-cjk-ui"' in text:
            issues.append(
                f"{locale}: CJK typography must only be applied to zh, zh-tw, ja and ko"
            )

    return issues


def check_tap_type_labels(generator, locales: dict[str, dict[str, str]]) -> list[str]:
    issues: list[str] = []
    targets = [
        ("English source", ENGLISH_APP, generator.TAP_TYPE_LABELS["en"]),
        ("main app", MAIN_APP, generator.TAP_TYPE_LABELS["en"]),
    ]
    targets.extend(
        (
            locale,
            app_path_for(locale, config),
            generator.TAP_TYPE_LABELS[locale],
        )
        for locale, config in locales.items()
    )

    js_keys = {
        "tt": "tapTypeTt",
        "tb": "tapTypeTb",
        "tfk": "tapTypeTfk",
    }
    for label, path, expected in targets:
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        for key, expected_label in expected.items():
            js_match = re.search(
                rf"{js_keys[key]}:\s*'([^']*)'",
                text,
            )
            actual_js_label = js_match.group(1) if js_match else None
            if actual_js_label != expected_label:
                issues.append(
                    f"{label}: expected JavaScript {key!r} tap-type label {expected_label!r}, found {actual_js_label!r}"
                )

            visible_match = re.search(
                rf'<input[^>]+name="scd-tap-type"[^>]+value="{re.escape(key)}"[^>]*>\s*<span>([^<]*)</span>',
                text,
            )
            actual_visible_label = html.unescape(visible_match.group(1)) if visible_match else None
            if actual_visible_label != expected_label:
                issues.append(
                    f"{label}: expected visible {key!r} tap-type label {expected_label!r}, found {actual_visible_label!r}"
                )
            if len(expected_label) > 46:
                issues.append(
                    f"{label}: tap-type label is too long ({len(expected_label)} characters): {expected_label!r}"
                )

    for label, path in (
        ("English source", ENGLISH_APP),
        ("main app", MAIN_APP),
    ):
        text = path.read_text(encoding="utf-8")
        for expected_label in ("with internal coolant", "without internal coolant"):
            if expected_label not in text:
                issues.append(f"{label}: missing normalized coolant label {expected_label!r}")
        if "with Internal Coolant" in text or "without Internal Coolant" in text:
            issues.append(f"{label}: coolant labels still contain inconsistent capitalization")

    return issues


def check_js_syntax(locales: dict[str, dict[str, str]]) -> list[str]:
    if not NODE.exists():
        return [f"Node runtime not found: {NODE}"]

    issues: list[str] = []
    with tempfile.TemporaryDirectory(prefix="smicalc-js-check-") as tmp:
        tmpdir = Path(tmp)
        for locale, config in locales.items():
            path = app_path_for(locale, config)
            if not path.exists():
                continue
            text = path.read_text(encoding="utf-8")
            scripts = re.findall(r"<script\b[^>]*>(.*?)</script>", text, flags=re.I | re.S)
            for index, script in enumerate(scripts, start=1):
                if not script.strip():
                    continue
                js_path = tmpdir / f"{locale}-{index}.js"
                js_path.write_text(script, encoding="utf-8")
                result = subprocess.run(
                    [str(NODE), "--check", str(js_path)],
                    text=True,
                    capture_output=True,
                    check=False,
                )
                if result.returncode:
                    message = (result.stderr or result.stdout).strip()
                    issues.append(f"{locale}: script block {index} failed node --check: {message}")
    return issues


def compile_false_friend_patterns() -> dict[str, list[re.Pattern[str]]]:
    case_sensitive_patterns = {
        r"EmpFehlung",
        r"等级",
        r"等級",
        r"등급",
    }
    return {
        group: [
            re.compile(pattern, 0 if pattern in case_sensitive_patterns else re.IGNORECASE)
            for pattern in patterns
        ]
        for group, patterns in FALSE_FRIEND_PATTERNS.items()
    }


def check_false_friends(locales: dict[str, dict[str, str]]) -> list[str]:
    patterns = compile_false_friend_patterns()
    issues: list[str] = []

    wb = load_workbook(WORKBOOK, read_only=True, data_only=False)
    ws = wb["Translations"]
    headers = [cell.value for cell in ws[1]]
    ignored_columns = {"English", "English (en-za_url)", "Swedish", "Context"}
    target_columns = {config["target_column"] for config in locales.values()}

    for row in ws.iter_rows(min_row=2):
        row_number = row[0].row
        for index, cell in enumerate(row, start=1):
            header = headers[index - 1]
            if header in ignored_columns or header not in target_columns:
                continue
            value = cell.value
            if not isinstance(value, str):
                continue
            for group, regexes in patterns.items():
                for regex in regexes:
                    if regex.search(value):
                        issues.append(
                            f"workbook row {row_number}, {header}: {group} matched {regex.pattern!r} in {value!r}"
                        )

    for locale, config in locales.items():
        path = app_path_for(locale, config)
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        for group, regexes in patterns.items():
            for regex in regexes:
                if regex.search(text):
                    issues.append(f"{locale}: {group} matched {regex.pattern!r} in {path.relative_to(ROOT)}")

    return issues


def check_grooving_localization(generator, locales: dict[str, dict[str, str]]) -> list[str]:
    """Protect navigation/DOM identifiers and the new warning placeholders."""
    issues: list[str] = []
    source = ENGLISH_APP.read_text(encoding="utf-8")
    source_arrays = re.findall(r"\[(?:\s*'[^']*'\s*,?)+\](?=\.forEach\(function\(key\))", source)
    source_placeholders = {
        key: sorted(re.findall(r"\{\w+\}", value))
        for key, value in re.findall(r"(grooving\w*):\s*'((?:\\.|[^'\\])*)'", source)
    }
    for locale, config in locales.items():
        path = app_path_for(locale, config)
        if not path.exists():
            continue
        content = path.read_text(encoding="utf-8")
        start = "        function parseGroovingNumber(value) {"
        end = "        function calculateThreadTurningData() {"
        if start not in content or end not in content or content.split(start, 1)[1].split(end, 1)[0] != source.split(start, 1)[1].split(end, 1)[0]:
            issues.append(f"{locale}: Grooving calculation code differs from English source")
        for token in ["section !== 'grooving'", *source_arrays]:
            if token not in content:
                issues.append(f"{locale}: Grooving machine identifier changed: {token[:80]}")
        values = dict(re.findall(r"(grooving\w*):\s*'((?:\\.|[^'\\])*)'", content))
        for key, expected in source_placeholders.items():
            if key not in values or sorted(re.findall(r"\{\w+\}", values[key])) != expected:
                issues.append(f"{locale}: missing Grooving text or changed placeholders: {key}")
        mapping = generator.load_translation_map(config["target_column"])
        for key in ("grooving", "groovingNotRecommended", "groovingWidthDepthWarning", "groovingDiameterWarning"):
            match = re.search(rf"\b{key}:\s*'((?:\\.|[^'\\])*)'", source)
            if match and match.group(1) not in mapping:
                issues.append(f"{locale}: missing workbook translation for {key}")
    return issues


def main() -> None:
    parser = argparse.ArgumentParser(description="Verify generated SmiCalc localized app HTML files.")
    parser.add_argument(
        "--locales",
        default="all",
        help="Comma-separated generator locale keys, or 'all'. Example: es,es-gl,da",
    )
    parser.add_argument("--skip-js", action="store_true", help="Skip JavaScript syntax checks.")
    parser.add_argument("--skip-api-language", action="store_true", help="Skip material selector API language checks.")
    parser.add_argument("--skip-false-friends", action="store_true", help="Skip targeted false-friend checks.")
    args = parser.parse_args()

    generator = load_generator()
    locales = selected_locales(generator.LOCALES, args.locales)
    issues: list[str] = []

    if not args.skip_api_language:
        issues.extend(check_api_language(locales))
    issues.extend(check_material_keys(locales))
    issues.extend(check_material_selector_history(locales))
    issues.extend(check_translation_workbook_canonical_labels(generator, locales))
    issues.extend(check_japanese_brand_names(locales))
    issues.extend(check_material_label_deduplication(locales))
    issues.extend(check_adjustable_input_accessibility(locales))
    issues.extend(check_cjk_typography(generator, locales))
    issues.extend(check_thread_profile_labels(generator, locales))
    issues.extend(check_tap_type_labels(generator, locales))
    issues.extend(check_grooving_localization(generator, locales))
    if not args.skip_js:
        issues.extend(check_js_syntax(locales))
    if not args.skip_false_friends:
        issues.extend(check_false_friends(locales))

    if issues:
        for issue in issues:
            print(f"FAIL: {issue}")
        raise SystemExit(1)

    print(f"OK: verified {len(locales)} locale(s)")


if __name__ == "__main__":
    main()
