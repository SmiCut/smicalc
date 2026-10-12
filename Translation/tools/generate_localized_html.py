from __future__ import annotations

import html
import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
TRANSLATION_WORKBOOK = ROOT / "outputs/translation-basis-20260603/smicut-cutting-data-translation-basis.xlsx"
URL_WORKBOOK = ROOT / "outputs/smicut_language_url_map_smicalc.xlsx"
MATERIAL_CATEGORY_KEYS_FILE = ROOT / "Translation/material-category-keys.json"
MATERIAL_CATEGORY_KEYS = json.loads(MATERIAL_CATEGORY_KEYS_FILE.read_text(encoding="utf-8"))

SOURCE_FILES = {
    "smicut cutting data app-eng.html": "smicut cutting data app-{suffix}.html",
    "smicut material categories-eng.html": "smicut material categories-{suffix}.html",
}

URL_SOURCE_ALIASES = {
    "https://smicut.com/en/knowledge/thread-turning/": "https://smicut.com/en/products/thread-turning-tools/",
}

LOCALES = {
    "sv": {
        "folder": "Svenska",
        "target_column": "Swedish",
        "url_column": "sv_url",
        "suffix": "se",
        "api_language": "sv",
        "source_folder": "English",
    },
    "es": {
        "folder": "Spanska",
        "target_column": "Spanish",
        "url_column": "es_url",
        "suffix": "es",
        "api_language": "es",
        "source_folder": "English",
    },
    "es-gl": {
        "folder": "Spanska Global",
        "target_column": "Spanish (es-gl_url)",
        "url_column": "es-gl_url",
        "suffix": "es-gl",
        "api_language": "es-gl",
        "source_folder": "English",
    },
    "id": {
        "folder": "Indonesiska",
        "target_column": "Indonesian (id_url)",
        "url_column": "id_url",
        "suffix": "id",
        "api_language": "id",
        "source_folder": "English",
    },
    "da": {
        "folder": "Danska",
        "target_column": "Danish (da_url)",
        "url_column": "da_url",
        "suffix": "da",
        "api_language": "da",
        "source_folder": "English",
    },
    "de": {
        "folder": "Tyska",
        "target_column": "German (de_url)",
        "url_column": "de_url",
        "suffix": "de",
        "api_language": "de",
        "source_folder": "English",
    },
    "en-za": {
        "folder": "English ZA",
        "target_column": "English (en-za_url)",
        "url_column": "en-za_url",
        "suffix": "en-za",
        "api_language": "en-za",
        "source_folder": "English",
    },
    "fr-be": {
        "folder": "Franska Belgien",
        "target_column": "French Belgium (fr-be_url)",
        "url_column": "fr-be_url",
        "suffix": "fr-be",
        "api_language": "fr-be",
        "source_folder": "English",
    },
    "fr": {
        "folder": "Franska",
        "target_column": "French (fr_url)",
        "url_column": "fr_url",
        "suffix": "fr",
        "api_language": "fr",
        "source_folder": "English",
    },
    "it": {
        "folder": "Italienska",
        "target_column": "Italian (it_url)",
        "url_column": "it_url",
        "suffix": "it",
        "api_language": "it",
        "source_folder": "English",
    },
    "hu": {
        "folder": "Ungerska",
        "target_column": "Hungarian (hu_url)",
        "url_column": "hu_url",
        "suffix": "hu",
        "api_language": "hu",
        "source_folder": "English",
    },
    "nl-be": {
        "folder": "Nederlandska Belgien",
        "target_column": "Dutch Belgium (nl-be_url)",
        "url_column": "nl-be_url",
        "suffix": "nl-be",
        "api_language": "nl-be",
        "source_folder": "English",
    },
    "nl": {
        "folder": "Nederlandska",
        "target_column": "Dutch (nl_url)",
        "url_column": "nl_url",
        "suffix": "nl",
        "api_language": "nl",
        "source_folder": "English",
    },
    "pl": {
        "folder": "Polska",
        "target_column": "Polish (pl_url)",
        "url_column": "pl_url",
        "suffix": "pl",
        "api_language": "pl",
        "source_folder": "English",
    },
    "ru": {
        "folder": "Ryska",
        "target_column": "Russian (ru_url)",
        "url_column": "ru_url",
        "suffix": "ru",
        "api_language": "ru",
        "source_folder": "English",
    },
    "zh": {
        "folder": "Kinesiska Simplified",
        "target_column": "Chinese Simplified (zh_url)",
        "url_column": "zh_url",
        "suffix": "zh",
        "api_language": "zh",
        "source_folder": "English",
    },
    "zh-tw": {
        "folder": "Kinesiska Traditional",
        "target_column": "Chinese Traditional (zh-tw_url)",
        "url_column": "zh-tw_url",
        "suffix": "zh-tw",
        "api_language": "zh-tw",
        "source_folder": "English",
    },
    "ja": {
        "folder": "Japanska",
        "target_column": "Japanese (ja_url)",
        "url_column": "ja_url",
        "suffix": "ja",
        "api_language": "ja",
        "source_folder": "English",
    },
    "ko": {
        "folder": "Koreanska",
        "target_column": "Korean (ko_url)",
        "url_column": "ko_url",
        "suffix": "ko",
        "api_language": "ko",
        "source_folder": "English",
    },
}

CJK_UI_LOCALES = {"zh", "zh-tw", "ja", "ko"}
CJK_TYPOGRAPHY_CSS = """
<style id="scd-cjk-typography">
    #smicut-cutting-data-app.scd-cjk-ui .scd-tab {
        font-weight: 600;
        white-space: nowrap;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-field-label {
        font-size: 16px;
        font-weight: 600;
        line-height: 1.35;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-control {
        font-size: 16px !important;
        font-weight: 500 !important;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-select,
    #smicut-cutting-data-app.scd-cjk-ui .scd-select option {
        font-weight: 500 !important;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-number {
        font-weight: 600 !important;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-choice span {
        font-size: 16px;
        font-weight: 600;
        line-height: 1.4;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-start-card strong,
    #smicut-cutting-data-app.scd-cjk-ui .scd-start-intro-title,
    #smicut-cutting-data-app.scd-cjk-ui .scd-section-title,
    #smicut-cutting-data-app.scd-cjk-ui .scd-note-title,
    #smicut-cutting-data-app.scd-cjk-ui .scd-result-title,
    #smicut-cutting-data-app.scd-cjk-ui #smicut-chat-widget .sms-header,
    #smicut-cutting-data-app.scd-cjk-ui #smicut-chat-widget .sms-reply-heading {
        font-weight: 700;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-start-card span {
        font-size: 15.5px;
        font-weight: 500;
        line-height: 1.55;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-note-text {
        font-size: 15px;
        font-weight: 500;
        line-height: 1.45;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-input-summary .scd-result-label {
        font-size: 14px;
        font-weight: 600;
        line-height: 1.35 !important;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-input-summary .scd-result-value {
        font-size: 16px;
        font-weight: 600;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-thread-summary-label {
        font-size: 13px;
        font-weight: 600;
        line-height: 1.25;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-thread-summary-value {
        font-size: 14px;
        font-weight: 700;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-result-row-highlight .scd-result-label {
        font-size: 16px !important;
        font-weight: 600;
        line-height: 1.35 !important;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-recommendation-note {
        font-size: 14px;
        font-weight: 600;
        line-height: 1.4;
    }

    #smicut-cutting-data-app.scd-cjk-ui #smicut-chat-widget .sms-message {
        font-size: 17px;
        font-weight: 500;
        line-height: 1.6;
    }

    #smicut-cutting-data-app.scd-cjk-ui #smicut-chat-widget .sms-category-result-label {
        font-weight: 600;
    }

    #smicut-cutting-data-app.scd-cjk-ui #smicut-chat-widget .sms-category-result-value,
    #smicut-cutting-data-app.scd-cjk-ui #smicut-chat-widget .sms-final-recommendation-result {
        font-weight: 800;
    }

    #smicut-cutting-data-app.scd-cjk-ui .scd-start-link,
    #smicut-cutting-data-app.scd-cjk-ui .scd-copy-button,
    #smicut-cutting-data-app.scd-cjk-ui #smicut-chat-widget #smicut-send,
    #smicut-cutting-data-app.scd-cjk-ui #smicut-chat-widget .sms-apply-material {
        font-weight: 700 !important;
    }

    @media (min-width: 760px) {
        #smicut-cutting-data-app.scd-cjk-ui .scd-tab {
            font-size: 16px;
            line-height: 1.25;
        }

        #smicut-cutting-data-app.scd-cjk-ui .scd-control {
            font-size: 17px !important;
        }
    }
</style>
""".strip()

TRANSLATABLE_ATTRS = ("aria-label", "placeholder", "title", "alt")
UNSAFE_LOWERCASE_JS_ALIASES = {"Internal", "External", "Triangular", "Grooving", "Diameter", "Copy"}

UI_TRANSLATION_OVERRIDES = {
    "id": {
        "without internal coolant": "Tanpa internal coolant",
        "with internal coolant": "Dengan internal coolant",
        "Drill length": "Panjang bor",
        "Thread Profile": "Profil ulir",
        "Thread profile": "Profil ulir",
        "Thread size": "Ukuran ulir",
        "Thread diameter": "Diameter ulir",
        "Thread diameter (mm)": "Diameter ulir (mm)",
        "Tap type": "Jenis tap",
        "Recommended drill diameter": "Diameter bor yang direkomendasikan",
        "Threading type": "Posisi ulir",
        "External": "Ulir luar",
        "Internal": "Ulir dalam",
        "Insert type": "Jenis insert",
        "Triangular": "Segitiga",
        "EC - General Tapping": "EC - Tapping umum",
        "PC - Harder Materials": "PC - Material lebih keras",
        "RC - Form Taps": "RC - Form taps",
    },
    # Japanese SmiCut product names follow the official Japanese website's
    # katakana convention in user-facing text. Legal names, URLs, email
    # addresses, HTML IDs, and code identifiers remain unchanged.
    "ja": {
        "Use in SmiCalc": "スミカルクで使用する",
        "SmiCut Material Selector": "スミカット マテリアルセレクター",
        "Enter a material designation and the tool will recommend the right SmiCut category.": "材料の指定を入力すると、ツールは適切なスミカットカテゴリを推奨します。",
        "➜ Learn more about SmiCut threading taps and grades": "➜ スミカットのタップと材種について詳しく見る",
    },
}

MATERIAL_DESCRIPTOR_PRESERVE_CASE_LOCALES = {"de"}
MATERIAL_DESCRIPTOR_FORCE_LOWER_PREFIXES = {
    "de": ("Legiert",),
}
THREAD_SUMMARY_LABELS = {
    "sv": ("Gänga", "Gängdiameter", "Stigning", "Invändig gänga", "Utvändig gänga"),
    "es": ("Rosca", "Diámetro de rosca", "Paso", "Rosca interior", "Rosca exterior"),
    "es-gl": ("Rosca", "Diámetro de rosca", "Paso", "Rosca interior", "Rosca exterior"),
    "id": ("Ulir", "Diameter ulir", "Pitch", "Ulir dalam", "Ulir luar"),
    "da": ("Gevind", "Gevinddiameter", "Stigning", "Indvendigt gevind", "Udvendigt gevind"),
    "de": ("Gewinde", "Gewindedurchmesser", "Steigung", "Innengewinde", "Außengewinde"),
    "en-za": ("Thread", "Thread diameter", "Pitch", "Internal thread", "External thread"),
    "fr-be": ("Filetage", "Diamètre du filetage", "Pas", "Filetage intérieur", "Filetage extérieur"),
    "fr": ("Filetage", "Diamètre du filetage", "Pas", "Filetage intérieur", "Filetage extérieur"),
    "it": ("Filettatura", "Diametro filettatura", "Passo", "Filettatura interna", "Filettatura esterna"),
    "hu": ("Menet", "Menetátmérő", "Menetemelkedés", "Belső menet", "Külső menet"),
    "nl-be": ("Schroefdraad", "Schroefdraaddiameter", "Spoed", "Inwendige schroefdraad", "Uitwendige schroefdraad"),
    "nl": ("Schroefdraad", "Schroefdraaddiameter", "Spoed", "Inwendige schroefdraad", "Uitwendige schroefdraad"),
    "pl": ("Gwint", "Średnica gwintu", "Skok", "Gwint wewnętrzny", "Gwint zewnętrzny"),
    "ru": ("Резьба", "Диаметр резьбы", "Шаг", "Внутренняя резьба", "Наружная резьба"),
    "zh": ("螺纹", "螺纹直径", "螺距", "内螺纹", "外螺纹"),
    "zh-tw": ("螺紋", "螺紋直徑", "螺距", "內螺紋", "外螺紋"),
    "ja": ("ねじ", "ねじ径", "ピッチ", "内ねじ", "外ねじ"),
    "ko": ("나사", "나사 직경", "피치", "내경 나사", "외경 나사"),
}
THREAD_PROFILE_LABELS = {
    "en": {
        "metricCoarse": "M - Metric coarse",
        "metricFine": "M - Metric fine",
        "unc": "UNC - Unified coarse",
        "unf": "UNF - Unified fine",
        "unef": "UNEF - Unified extra fine",
        "unConstant": "UN - Unified constant pitch",
        "g": "G - Whitworth Pipe Thread",
        "gOther": "G - Whitworth Pipe Thread (manual)",
        "bspt": "BSPT - Tapered Pipe Thread",
        "pg": "PG - Steel Conduit Thread DIN 40430",
        "npt": "NPT - Tapered Pipe Thread",
        "nptf": "NPTF - Dryseal, Tapered Pipe Thread",
    },
    "sv": {
        "metricCoarse": "M - Metrisk grov",
        "metricFine": "M - Metrisk fin",
        "unc": "UNC - Unified grov",
        "unf": "UNF - Unified fin",
        "unef": "UNEF - Unified extra fin",
        "unConstant": "UN - Konstant stigning",
        "g": "G - Whitworth rörgänga",
        "gOther": "G - Whitworth rörgänga (manuell)",
        "bspt": "BSPT - Konisk Whitworth rörgänga",
        "pg": "PG - Pansarrörsgänga DIN 40430",
        "npt": "NPT - Konisk rörgänga",
        "nptf": "NPTF - Dryseal, konisk rörgänga",
    },
    "es": {
        "metricCoarse": "M - Métrica normal",
        "metricFine": "M - Métrica fina",
        "unc": "UNC - Unified, paso grueso",
        "unf": "UNF - Unified, paso fino",
        "unef": "UNEF - Unified, paso extrafino",
        "unConstant": "UN - Paso constante",
        "g": "G - Rosca Whitworth para tubos",
        "gOther": "G - Rosca Whitworth para tubos (manual)",
        "bspt": "BSPT - Rosca Whitworth cónica para tubos",
        "pg": "PG - Panzerrohrgewinde DIN 40430",
        "npt": "NPT - Rosca cónica para tubos",
        "nptf": "NPTF - Rosca cónica Dryseal para tubos",
    },
    "es-gl": {
        "metricCoarse": "M - Métrica estándar",
        "metricFine": "M - Métrica fina",
        "unc": "UNC - Unified, paso grueso",
        "unf": "UNF - Unified, paso fino",
        "unef": "UNEF - Unified, paso extrafino",
        "unConstant": "UN - Paso constante",
        "g": "G - Rosca de tubo Whitworth",
        "gOther": "G - Rosca de tubo Whitworth (manual)",
        "bspt": "BSPT - Rosca Whitworth cónica de tubo",
        "pg": "PG - Panzerrohrgewinde DIN 40430",
        "npt": "NPT - Rosca cónica de tubo",
        "nptf": "NPTF - Rosca cónica Dryseal de tubo",
    },
    "id": {
        "metricCoarse": "M - Metrik kasar",
        "metricFine": "M - Metrik halus",
        "unc": "UNC - Unified kasar",
        "unf": "UNF - Unified halus",
        "unef": "UNEF - Unified ekstra halus",
        "unConstant": "UN - Pitch konstan",
        "g": "G - Ulir pipa Whitworth",
        "gOther": "G - Ulir pipa Whitworth (manual)",
        "bspt": "BSPT - Ulir pipa tirus Whitworth",
        "pg": "PG - Steel Conduit Thread DIN 40430",
        "npt": "NPT - Ulir pipa tirus",
        "nptf": "NPTF - Dryseal, ulir pipa tirus",
    },
    "da": {
        "metricCoarse": "M - Metrisk grov",
        "metricFine": "M - Metrisk fin",
        "unc": "UNC - Unified grov",
        "unf": "UNF - Unified fin",
        "unef": "UNEF - Unified ekstra fin",
        "unConstant": "UN - Konstant stigning",
        "g": "G - Whitworth-rørgevind",
        "gOther": "G - Whitworth-rørgevind (manuelt)",
        "bspt": "BSPT - Konisk Whitworth-rørgevind",
        "pg": "PG - Panserrørgevind DIN 40430",
        "npt": "NPT - Konisk rørgevind",
        "nptf": "NPTF - Dryseal, konisk rørgevind",
    },
    "de": {
        "metricCoarse": "M - Metrisches Regelgewinde",
        "metricFine": "M - Metrisches Feingewinde",
        "unc": "UNC - Unified-Grobgewinde",
        "unf": "UNF - Unified-Feingewinde",
        "unef": "UNEF - Unified-Extra-Feingewinde",
        "unConstant": "UN - Unified-Gewinde mit konstanter Steigung",
        "g": "G - Whitworth-Rohrgewinde",
        "gOther": "G - Whitworth-Rohrgewinde (manuell)",
        "bspt": "BSPT - Kegeliges Whitworth-Rohrgewinde",
        "pg": "PG - Panzerrohrgewinde DIN 40430",
        "npt": "NPT - Kegeliges Rohrgewinde",
        "nptf": "NPTF - Dryseal, kegeliges Rohrgewinde",
    },
    "en-za": {
        "metricCoarse": "M - Metric coarse",
        "metricFine": "M - Metric fine",
        "unc": "UNC - Unified coarse",
        "unf": "UNF - Unified fine",
        "unef": "UNEF - Unified extra fine",
        "unConstant": "UN - Unified constant pitch",
        "g": "G - Whitworth Pipe Thread",
        "gOther": "G - Whitworth Pipe Thread (manual)",
        "bspt": "BSPT - Tapered Pipe Thread",
        "pg": "PG - Steel Conduit Thread DIN 40430",
        "npt": "NPT - Tapered Pipe Thread",
        "nptf": "NPTF - Dryseal, Tapered Pipe Thread",
    },
    "fr-be": {
        "metricCoarse": "M - Métrique, gros pas",
        "metricFine": "M - Métrique, pas fin",
        "unc": "UNC - Gros pas",
        "unf": "UNF - Pas fin",
        "unef": "UNEF - Pas extra-fin",
        "unConstant": "UN - Pas constant",
        "g": "G - Gaz cylindrique Whitworth",
        "gOther": "G - Gaz cylindrique Whitworth (manuel)",
        "bspt": "BSPT - Gaz conique Whitworth",
        "pg": "PG - DIN 40430",
        "npt": "NPT - Conique pour tubes",
        "nptf": "NPTF - Dryseal conique pour tubes",
    },
    "fr": {
        "metricCoarse": "M - Métrique, gros pas",
        "metricFine": "M - Métrique, pas fin",
        "unc": "UNC - Gros pas",
        "unf": "UNF - Pas fin",
        "unef": "UNEF - Pas extra-fin",
        "unConstant": "UN - Pas constant",
        "g": "G - Gaz cylindrique Whitworth",
        "gOther": "G - Gaz cylindrique Whitworth (manuel)",
        "bspt": "BSPT - Gaz conique Whitworth",
        "pg": "PG - DIN 40430",
        "npt": "NPT - Conique pour tubes",
        "nptf": "NPTF - Dryseal conique pour tubes",
    },
    "it": {
        "metricCoarse": "M - Metrica, passo grosso",
        "metricFine": "M - Metrica, passo fine",
        "unc": "UNC - Unified, passo grosso",
        "unf": "UNF - Unified, passo fine",
        "unef": "UNEF - Unified, passo extra fine",
        "unConstant": "UN - Passo costante",
        "g": "G - Gas cilindrica Whitworth",
        "gOther": "G - Gas cilindrica Whitworth (manuale)",
        "bspt": "BSPT - Gas conica Whitworth",
        "pg": "PG - Filettatura per tubi elettrici corazzati DIN 40430",
        "npt": "NPT - Conica per tubi",
        "nptf": "NPTF - Dryseal conica per tubi",
    },
    "hu": {
        "metricCoarse": "M - Metrikus, durva",
        "metricFine": "M - Metrikus, finom",
        "unc": "UNC - Unified, durva",
        "unf": "UNF - Unified, finom",
        "unef": "UNEF - Unified, extra finom",
        "unConstant": "UN - Állandó emelkedés",
        "g": "G - Whitworth csőmenet",
        "gOther": "G - Whitworth csőmenet (kézi)",
        "bspt": "BSPT - Kúpos Whitworth csőmenet",
        "pg": "PG - Páncélcsőmenet DIN 40430",
        "npt": "NPT - Kúpos csőmenet",
        "nptf": "NPTF - Dryseal kúpos csőmenet",
    },
    "nl-be": {
        "metricCoarse": "M - Metrisch grof",
        "metricFine": "M - Metrisch fijn",
        "unc": "UNC - Unified grof",
        "unf": "UNF - Unified fijn",
        "unef": "UNEF - Unified extra fijn",
        "unConstant": "UN - Constante spoed",
        "g": "G - Whitworth-pijpdraad",
        "gOther": "G - Whitworth-pijpdraad (handmatig)",
        "bspt": "BSPT - Conische Whitworth-pijpdraad",
        "pg": "PG - Panzerrohrgewinde DIN 40430",
        "npt": "NPT - Conische pijpdraad",
        "nptf": "NPTF - Dryseal, conische pijpdraad",
    },
    "nl": {
        "metricCoarse": "M - Metrisch grof",
        "metricFine": "M - Metrisch fijn",
        "unc": "UNC - Unified grof",
        "unf": "UNF - Unified fijn",
        "unef": "UNEF - Unified extra fijn",
        "unConstant": "UN - Constante spoed",
        "g": "G - Whitworth-pijpdraad",
        "gOther": "G - Whitworth-pijpdraad (handmatig)",
        "bspt": "BSPT - Conische Whitworth-pijpdraad",
        "pg": "PG - Panzerrohrgewinde DIN 40430",
        "npt": "NPT - Conische pijpdraad",
        "nptf": "NPTF - Dryseal, conische pijpdraad",
    },
    "pl": {
        "metricCoarse": "M - metryczny zwykły",
        "metricFine": "M - metryczny drobny",
        "unc": "UNC - zunifikowany gruby",
        "unf": "UNF - zunifikowany drobny",
        "unef": "UNEF - zunifikowany bardzo drobny",
        "unConstant": "UN - stały skok",
        "g": "G - gwint rurowy Whitwortha",
        "gOther": "G - gwint rurowy Whitwortha (ręcznie)",
        "bspt": "BSPT - stożkowy gwint rurowy Whitwortha",
        "pg": "PG - gwint do rurek instalacyjnych DIN 40430",
        "npt": "NPT - gwint rurowy stożkowy",
        "nptf": "NPTF - Dryseal, gwint rurowy stożkowy",
    },
    "ru": {
        "metricCoarse": "M - метрическая, крупная",
        "metricFine": "M - метрическая, мелкая",
        "unc": "UNC - унифицированная, крупная",
        "unf": "UNF - унифицированная, мелкая",
        "unef": "UNEF - унифицированная, особо мелкая",
        "unConstant": "UN - постоянный шаг",
        "g": "G - трубная резьба Витворта",
        "gOther": "G - трубная резьба Витворта (вручную)",
        "bspt": "BSPT - коническая трубная резьба Витворта",
        "pg": "PG - резьба электротехнического назначения DIN 40430",
        "npt": "NPT - коническая трубная резьба",
        "nptf": "NPTF - коническая трубная резьба Dryseal",
    },
    "zh": {
        "metricCoarse": "M - 公制粗牙",
        "metricFine": "M - 公制细牙",
        "unc": "UNC - 统一粗牙",
        "unf": "UNF - 统一细牙",
        "unef": "UNEF - 统一超细牙",
        "unConstant": "UN - 恒定螺距",
        "g": "G - 惠氏管螺纹",
        "gOther": "G - 惠氏管螺纹（手动）",
        "bspt": "BSPT - 英制锥管螺纹",
        "pg": "PG - DIN 40430 钢管螺纹",
        "npt": "NPT - 锥管螺纹",
        "nptf": "NPTF - 干密封锥管螺纹",
    },
    "zh-tw": {
        "metricCoarse": "M - 公制粗牙",
        "metricFine": "M - 公制細牙",
        "unc": "UNC - 統一粗牙",
        "unf": "UNF - 統一細牙",
        "unef": "UNEF - 統一特細牙",
        "unConstant": "UN - 恆定螺距",
        "g": "G - 英制直管螺紋",
        "gOther": "G - 英制直管螺紋（手動）",
        "bspt": "BSPT - 英制錐管螺紋",
        "pg": "PG - 鋼管牙 DIN 40430",
        "npt": "NPT - 錐管螺紋",
        "nptf": "NPTF - 乾密封錐管螺紋",
    },
    "ja": {
        "metricCoarse": "M - メートル並目",
        "metricFine": "M - メートル細目",
        "unc": "UNC - ユニファイ並目",
        "unf": "UNF - ユニファイ細目",
        "unef": "UNEF - ユニファイ極細目",
        "unConstant": "UN - ユニファイ（一定ピッチ）",
        "g": "G - 管用平行ねじ",
        "gOther": "G - 管用平行ねじ（手動）",
        "bspt": "BSPT - 管用テーパねじ",
        "pg": "PG - ＰＧねじ DIN 40430",
        "npt": "NPT - 管用テーパねじ",
        "nptf": "NPTF - ドライシール管用テーパねじ",
    },
    "ko": {
        "metricCoarse": "M - 미터 보통",
        "metricFine": "M - 미터 가는",
        "unc": "UNC - 유니파이 보통",
        "unf": "UNF - 유니파이 가는",
        "unef": "UNEF - 유니파이 매우 가는",
        "unConstant": "UN - 유니파이 일정 피치",
        "g": "G - 관용 위트워드 나사",
        "gOther": "G - 관용 위트워드 나사(수동)",
        "bspt": "BSPT - 관용 테이퍼 나사(영국식)",
        "pg": "PG - 전선관용 나사 DIN 40430",
        "npt": "NPT - 관용 테이퍼 나사",
        "nptf": "NPTF - 드라이씰용 관용 테이퍼 나사",
    },
}
TAP_TYPE_LABELS = {
    "en": {
        "tt": "TT - Spiral point, through holes",
        "tb": "TB - Spiral flute, blind holes",
        "tfk": "TFK - Forming",
    },
    "sv": {
        "tt": "TT - Spåndrivande, genomgående hål",
        "tb": "TB - Spiralskuren, bottenhål",
        "tfk": "TFK - Formning",
    },
    "es": {
        "tt": "TT - Punta helicoidal, agujero pasante",
        "tb": "TB - Canal helicoidal, agujero ciego",
        "tfk": "TFK - Laminación",
    },
    "es-gl": {
        "tt": "TT - Punta helicoidal, agujero pasante",
        "tb": "TB - Canal helicoidal, agujero ciego",
        "tfk": "TFK - Laminación",
    },
    "id": {
        "tt": "TT - Spiral point, lubang tembus",
        "tb": "TB - Spiral flute, lubang buntu",
        "tfk": "TFK - Forming",
    },
    "da": {
        "tt": "TT - Spiralspids, gennemgående hul",
        "tb": "TB - Spiralrille, bundhul",
        "tfk": "TFK - Formning",
    },
    "de": {
        "tt": "TT - Schälanschnitt, Durchgangsloch",
        "tb": "TB - Spiralnut, Sackloch",
        "tfk": "TFK - Gewindeformen",
    },
    "en-za": {
        "tt": "TT - Spiral point, through holes",
        "tb": "TB - Spiral flute, blind holes",
        "tfk": "TFK - Forming",
    },
    "fr-be": {
        "tt": "TT - Entrée hélicoïdale, trou débouchant",
        "tb": "TB - Goujures hélicoïdales, trou borgne",
        "tfk": "TFK - Déformation",
    },
    "fr": {
        "tt": "TT - Entrée hélicoïdale, trou débouchant",
        "tb": "TB - Goujures hélicoïdales, trou borgne",
        "tfk": "TFK - Déformation",
    },
    "it": {
        "tt": "TT - Punta elicoidale, foro passante",
        "tb": "TB - Scanalatura elicoidale, foro cieco",
        "tfk": "TFK - Rullatura",
    },
    "hu": {
        "tt": "TT - Spirálhegy, átmenő furat",
        "tb": "TB - Spirálhorony, zsákfurat",
        "tfk": "TFK - Menetformázás",
    },
    "nl-be": {
        "tt": "TT - Spiraalpunt, doorlopend gat",
        "tb": "TB - Spiraalgroef, blind gat",
        "tfk": "TFK - Draadvormen",
    },
    "nl": {
        "tt": "TT - Spiraalpunt, doorlopend gat",
        "tb": "TB - Spiraalgroef, blind gat",
        "tfk": "TFK - Draadvormen",
    },
    "pl": {
        "tt": "TT - nakrój spiralny, otwór przelotowy",
        "tb": "TB - rowek spiralny, otwór nieprzelotowy",
        "tfk": "TFK - formowanie",
    },
    "ru": {
        "tt": "TT - подточка, сквозное отверстие",
        "tb": "TB - спиральная канавка, глухое отверстие",
        "tfk": "TFK - раскатывание",
    },
    "zh": {
        "tt": "TT - 螺旋尖，通孔",
        "tb": "TB - 螺旋槽，盲孔",
        "tfk": "TFK - 挤压",
    },
    "zh-tw": {
        "tt": "TT - 先端，通孔",
        "tb": "TB - 螺旋槽，盲孔",
        "tfk": "TFK - 成型",
    },
    "ja": {
        "tt": "TT - ポイント・通り穴",
        "tb": "TB - スパイラル・止まり穴",
        "tfk": "TFK - 転造",
    },
    "ko": {
        "tt": "TT - 스파이럴 포인트, 관통 홀",
        "tb": "TB - 스파이럴 플루트, 막힌 홀",
        "tfk": "TFK - 성형(전조)",
    },
}
START_PAGE_LABELS = {
    "sv": {
        "smiprog": "SmiProg - Gängfräsning",
        "milling_title": "Fräsning",
        "milling_desc": "Beräkna skärhastighet, matning per tand, spindelvarvtal och bordsmatning för solida hårdmetallfräsar baserat på diameter, antal skär, ap och ae.",
        "drilling_title": "Borrning",
        "drilling_desc": "Beräkna skärhastighet, matning per varv, spindelvarvtal och bordsmatning för solida hårdmetallborr med valfria egna värden.",
    },
    "es": {
        "smiprog": "SmiProg - Fresado de roscas",
        "milling_title": "Fresado",
        "milling_desc": "Calcule la velocidad de corte, el avance por labio, la velocidad del husillo y el avance de mesa para fresas de metal duro integral utilizando el diámetro, el número de labios, ap y ae.",
        "drilling_title": "Taladrado",
        "drilling_desc": "Calcule la velocidad de corte, el avance por revolución, la velocidad del husillo y el avance de mesa para brocas de metal duro integral con valores propios opcionales.",
    },
    "es-gl": {
        "smiprog": "SmiProg - Fresado de roscas",
        "milling_title": "Fresado",
        "milling_desc": "Calcule la velocidad de corte, el avance por filo, la velocidad del husillo y el avance de mesa para fresas de carburo sólido utilizando el diámetro, el número de filos, ap y ae.",
        "drilling_title": "Taladrado",
        "drilling_desc": "Calcule la velocidad de corte, el avance por revolución, la velocidad del husillo y el avance de mesa para brocas de carburo sólido con valores propios opcionales.",
    },
    "id": {
        "smiprog": "SmiProg - Thread Milling",
        "milling_title": "Milling",
        "milling_desc": "Hitung cutting speed, feed per flute, spindle speed, dan table feed untuk solid carbide end mills menggunakan diameter, jumlah flute, ap, dan ae.",
        "drilling_title": "Drilling",
        "drilling_desc": "Hitung cutting speed, feed per revolution, spindle speed, dan table feed untuk solid carbide drills dengan opsi nilai sendiri.",
    },
    "da": {
        "smiprog": "SmiProg - Gevindfræsning",
        "milling_title": "Fræsning",
        "milling_desc": "Beregn skærehastighed, tilspænding pr. tand, spindelomdrejningstal og bordtilspænding for pindfræsere i solid hårdmetal baseret på diameter, antal skær, ap og ae.",
        "drilling_title": "Boring",
        "drilling_desc": "Beregn skærehastighed, tilspænding pr. omdrejning, spindelomdrejningstal og bordtilspænding for bor i solid hårdmetal med valgfrie egne værdier.",
    },
    "de": {
        "smiprog": "SmiProg - Gewindefräsen",
        "milling_title": "Fräsen",
        "milling_desc": "Berechnen Sie Schnittgeschwindigkeit, Vorschub pro Zahn, Spindeldrehzahl und Tischvorschub für Vollhartmetall-Schaftfräser anhand von Durchmesser, Schneidenzahl, ap und ae.",
        "drilling_title": "Bohren",
        "drilling_desc": "Berechnen Sie Schnittgeschwindigkeit, Vorschub pro Umdrehung, Spindeldrehzahl und Tischvorschub für Vollhartmetallbohrer mit optionalen eigenen Werten.",
    },
    "en-za": {
        "smiprog": "SmiProg - Thread Milling",
        "milling_title": "Milling",
        "milling_desc": "Calculate cutting speed, feed per flute, spindle speed and table feed for solid carbide end mills using diameter, flutes, ap and ae.",
        "drilling_title": "Drilling",
        "drilling_desc": "Calculate cutting speed, feed per revolution, spindle speed and table feed for solid carbide drills with optional own values.",
    },
    "fr-be": {
        "smiprog": "SmiProg - Fraisage de filets",
        "milling_title": "Fraisage",
        "milling_desc": "Calculez la vitesse de coupe, l'avance par dent, la vitesse de broche et l'avance de table pour les fraises en carbure monobloc à partir du diamètre, du nombre de dents et des valeurs ap et ae.",
        "drilling_title": "Perçage",
        "drilling_desc": "Calculez la vitesse de coupe, l'avance par tour, la vitesse de broche et l'avance de table pour les forets en carbure monobloc avec vos propres valeurs en option.",
    },
    "fr": {
        "smiprog": "SmiProg - Fraisage de filets",
        "milling_title": "Fraisage",
        "milling_desc": "Calculez la vitesse de coupe, l'avance par dent, la vitesse de broche et l'avance de table pour les fraises en carbure monobloc à partir du diamètre, du nombre de dents et des valeurs ap et ae.",
        "drilling_title": "Perçage",
        "drilling_desc": "Calculez la vitesse de coupe, l'avance par tour, la vitesse de broche et l'avance de table pour les forets en carbure monobloc avec vos propres valeurs en option.",
    },
    "it": {
        "smiprog": "SmiProg - Fresatura dei filetti",
        "milling_title": "Fresatura",
        "milling_desc": "Calcola la velocità di taglio, l'avanzamento per dente, la velocità del mandrino e l'avanzamento tavola per frese in metallo duro integrale utilizzando diametro, numero di taglienti, ap e ae.",
        "drilling_title": "Foratura",
        "drilling_desc": "Calcola la velocità di taglio, l'avanzamento per giro, la velocità del mandrino e l'avanzamento tavola per punte in metallo duro integrale con valori personalizzati opzionali.",
    },
    "hu": {
        "smiprog": "SmiProg - Menetmarás",
        "milling_title": "Marás",
        "milling_desc": "Számítsa ki a forgácsolási sebességet, a fogankénti előtolást, az orsófordulatszámot és az asztalelőtolást keményfém ujjmarókhoz az átmérő, az élek száma, az ap és az ae alapján.",
        "drilling_title": "Fúrás",
        "drilling_desc": "Számítsa ki a forgácsolási sebességet, a fordulatonkénti előtolást, az orsófordulatszámot és az asztalelőtolást keményfém csigafúrókhoz, opcionálisan saját értékekkel.",
    },
    "nl-be": {
        "smiprog": "SmiProg - Draadfrezen",
        "milling_title": "Frezen",
        "milling_desc": "Bereken snijsnelheid, voeding per tand, spiltoerental en tafelvoeding voor volhardmetalen vingerfrezen op basis van diameter, aantal tanden, ap en ae.",
        "drilling_title": "Boren",
        "drilling_desc": "Bereken snijsnelheid, voeding per omwenteling, spiltoerental en tafelvoeding voor volhardmetalen boren met optionele eigen waarden.",
    },
    "nl": {
        "smiprog": "SmiProg - Draadfrezen",
        "milling_title": "Frezen",
        "milling_desc": "Bereken snijsnelheid, voeding per tand, spiltoerental en tafelvoeding voor volhardmetalen vingerfrezen op basis van diameter, aantal snijkanten, ap en ae.",
        "drilling_title": "Boren",
        "drilling_desc": "Bereken snijsnelheid, voeding per omwenteling, spiltoerental en tafelvoeding voor volhardmetalen boren met optionele eigen waarden.",
    },
    "pl": {
        "smiprog": "SmiProg - Frezowanie gwintów",
        "milling_title": "Frezowanie",
        "milling_desc": "Oblicz prędkość skrawania, posuw na ostrze, prędkość obrotową wrzeciona i posuw stołu dla frezów z węglików spiekanych na podstawie średnicy, liczby ostrzy, ap i ae.",
        "drilling_title": "Wiercenie",
        "drilling_desc": "Oblicz prędkość skrawania, posuw na obrót, prędkość obrotową wrzeciona i posuw stołu dla wierteł z węglików spiekanych z opcjonalnymi własnymi wartościami.",
    },
    "ru": {
        "smiprog": "SmiProg - Фрезерование резьбы",
        "milling_title": "Фрезерование",
        "milling_desc": "Рассчитайте скорость резания, подачу на зуб, частоту вращения шпинделя и подачу стола для монолитных твердосплавных фрез на основе диаметра, числа зубьев, ap и ae.",
        "drilling_title": "Сверление",
        "drilling_desc": "Рассчитайте скорость резания, подачу на оборот, частоту вращения шпинделя и подачу стола для монолитных твердосплавных сверл с дополнительными собственными значениями.",
    },
    "zh": {
        "smiprog": "SmiProg - 螺纹铣削",
        "milling_title": "铣削",
        "milling_desc": "根据直径、刃数、ap 和 ae，为整体硬质合金立铣刀计算切削速度、每齿进给量、主轴转速和工作台进给量。",
        "drilling_title": "钻削",
        "drilling_desc": "使用可选的自定义值，为整体硬质合金钻头计算切削速度、每转进给量、主轴转速和工作台进给量。",
    },
    "zh-tw": {
        "smiprog": "SmiProg - 螺紋銑削",
        "milling_title": "銑削",
        "milling_desc": "根據直徑、刃數、ap 和 ae，為鎢鋼銑刀計算切削速度、每刃進給量、主軸轉速和工作台進給。",
        "drilling_title": "鑽削",
        "drilling_desc": "使用可選的自訂值，為鎢鋼鑽頭計算切削速度、每轉進給量、主軸轉速和工作台進給。",
    },
    "ja": {
        "smiprog": "スミプログ - ねじフライス加工",
        "milling_title": "フライス加工",
        "milling_desc": "超硬ソリッドエンドミルの切削速度、刃当たり送り、主軸回転数、テーブル送りを、工具径、刃数、ap、aeに基づいて計算します。",
        "drilling_title": "穴あけ",
        "drilling_desc": "超硬ソリッドドリルの切削速度、1回転当たり送り、主軸回転数、テーブル送りを、任意の独自値で計算します。",
    },
    "ko": {
        "smiprog": "SmiProg - 나사 밀링",
        "milling_title": "밀링",
        "milling_desc": "초경 엔드밀용 절삭 속도, 날당 이송, 스핀들 속도 및 테이블 이송을 직경, 날 수, ap 및 ae를 기준으로 계산합니다.",
        "drilling_title": "드릴링",
        "drilling_desc": "초경 드릴용 절삭 속도, 회전당 이송, 스핀들 속도 및 테이블 이송을 필요 시 직접 입력한 값으로 계산합니다.",
    },
}


def lower_first_character(value: str) -> str:
    for index, char in enumerate(value):
        if char.isspace():
            continue
        if char.isupper():
            return value[:index] + char.lower() + value[index + 1 :]
        return value
    return value


def add_translation_mapping(mapping: dict[str, str], source: str, target: str, *, add_lower_alias: bool = True) -> None:
    mapping[source] = target
    if not add_lower_alias or source == target:
        return
    lower_source = lower_first_character(source)
    if lower_source != source:
        mapping.setdefault(lower_source, target)


def load_translation_map(target_column: str) -> dict[str, str]:
    wb = load_workbook(TRANSLATION_WORKBOOK, data_only=False)
    ws = wb["Translations"]
    headers = [cell.value for cell in ws[1]]
    english_col = headers.index("English") + 1
    target_col = headers.index(target_column) + 1
    context_col = headers.index("Context") + 1 if "Context" in headers else None
    mapping: dict[str, str] = {}

    for row in range(2, ws.max_row + 1):
        english_value = ws.cell(row, english_col).value
        target_value = ws.cell(row, target_col).value
        context_value = ws.cell(row, context_col).value if context_col else ""
        if english_value is None or target_value is None:
            continue
        english_text = str(english_value)
        target_text = str(target_value)
        add_translation_mapping(
            mapping,
            english_text,
            target_text,
            add_lower_alias=english_text not in UNSAFE_LOWERCASE_JS_ALIASES,
        )
        if english_text.endswith(":") and not target_text.endswith(" "):
            add_translation_mapping(mapping, f"{english_text} ", f"{target_text} ")

        # Material category HTML uses comma decimals, while the JS data uses dot decimals.
        decimal_variant = re.sub(r"(?<=\d)\.(?=\d)", ",", english_text)
        if decimal_variant != english_text:
            add_translation_mapping(mapping, decimal_variant, target_text)

        # Result-value unit suffixes are stored without leading whitespace in the workbook,
        # while the JavaScript source appends them as string literals such as " mm/flute".
        if "unit suffix" in str(context_value).lower():
            add_translation_mapping(mapping, f" {english_text}", f" {target_text}")

    return mapping


def load_url_map(url_column: str) -> dict[str, str]:
    if not URL_WORKBOOK.exists():
        return {}

    wb = load_workbook(URL_WORKBOOK, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    if "en_url" not in headers or url_column not in headers:
        return {}

    en_col = headers.index("en_url") + 1
    target_col = headers.index(url_column) + 1
    mapping: dict[str, str] = {}
    for row in range(2, ws.max_row + 1):
        source = ws.cell(row, en_col).value
        target = ws.cell(row, target_col).value
        if source and target:
            mapping[str(source)] = str(target)
    for alias_source, canonical_source in URL_SOURCE_ALIASES.items():
        canonical_target = mapping.get(canonical_source)
        if canonical_target:
            mapping[alias_source] = canonical_target
    return mapping


def js_escape(value: str) -> str:
    return (
        value.replace("\\", "\\\\")
        .replace("\n", "\\n")
        .replace("\r", "\\r")
        .replace("'", "\\'")
    )


def translate_text(value: str, mapping: dict[str, str]) -> str:
    return mapping.get(value, value)


def split_for_br(value: str) -> str:
    unit_match = re.search(r"\s+(HB|N/mm²)$", value)
    if unit_match:
        return value[: unit_match.start()] + "<br>" + unit_match.group(1)
    if value == "Inicio":
        return value
    words = value.split()
    if len(words) >= 3:
        midpoint = (len(words) + 1) // 2
        return " ".join(words[:midpoint]) + "<br>" + " ".join(words[midpoint:])
    if len(words) == 2:
        return words[0] + "<br>" + words[1]
    return value


def replace_br_sequences(content: str, mapping: dict[str, str]) -> str:
    def repl(match: re.Match[str]) -> str:
        first = html.unescape(match.group("first")).strip()
        second = html.unescape(match.group("second")).strip()
        source = f"{first} {second}"
        target = mapping.get(source)
        if not target:
            return match.group(0)
        return split_for_br(target)

    return re.sub(
        r"(?P<first>[^<>\n]+)<br\s*/?>(?P<second>[^<>\n]+)",
        repl,
        content,
        flags=re.IGNORECASE,
    )


def replace_attrs(content: str, mapping: dict[str, str]) -> str:
    attrs = "|".join(re.escape(attr) for attr in TRANSLATABLE_ATTRS)
    pattern = re.compile(rf"\b(?P<attr>{attrs})=(?P<quote>[\"'])(?P<value>.*?)(?P=quote)")

    def repl(match: re.Match[str]) -> str:
        value = html.unescape(match.group("value"))
        target = mapping.get(value)
        if not target:
            return match.group(0)
        escaped = html.escape(target, quote=True)
        return f"{match.group('attr')}={match.group('quote')}{escaped}{match.group('quote')}"

    return pattern.sub(repl, content)


def replace_text_nodes(content: str, mapping: dict[str, str]) -> str:
    text_pattern = re.compile(r">(?P<text>[^<>]+)<")

    def repl(match: re.Match[str]) -> str:
        raw = match.group("text")
        if not raw.strip():
            return match.group(0)
        leading = raw[: len(raw) - len(raw.lstrip())]
        trailing = raw[len(raw.rstrip()) :]
        core = raw.strip()
        target = mapping.get(html.unescape(core))
        if not target:
            return match.group(0)
        return ">" + leading + html.escape(target, quote=False) + trailing + "<"

    return text_pattern.sub(repl, content)


def replace_js_strings(script: str, mapping: dict[str, str]) -> str:
    string_pattern = re.compile(r"(?P<quote>['\"])(?P<value>(?:\\.|(?!\\1).)*?)(?P=quote)")
    protected_key_values: list[str] = []

    def protect_key_value(match: re.Match[str]) -> str:
        protected_key_values.append(match.group(0))
        return f"@@SMICUT_MACHINE_KEY_{len(protected_key_values) - 1}@@"

    # Material/category keys are API identifiers. They must never be translated,
    # even when their visible English value also appears in the translation map.
    script = re.sub(
        r"\bkey\s*:\s*(?:'(?:\\.|[^'\\])*'|\"(?:\\.|[^\"\\])*\")",
        protect_key_value,
        script,
    )
    # Grooving builds DOM references from suffix arrays. These are identifiers,
    # including words such as material/copy/diameter that also label the UI.
    script = re.sub(
        r"\[(?:\s*'[^']*'\s*,?)+\](?=\.forEach\(function\(key\))",
        protect_key_value,
        script,
    )

    def unescape_js(value: str) -> str:
        return (
            value.replace("\\'", "'")
            .replace('\\"', '"')
            .replace("\\n", "\n")
            .replace("\\r", "\r")
            .replace("\\\\", "\\")
        )

    def repl(match: re.Match[str]) -> str:
        source = unescape_js(match.group("value"))
        target = mapping.get(source)
        if not target:
            return match.group(0)
        quote = match.group("quote")
        if quote == "'":
            return "'" + js_escape(target) + "'"
        return '"' + target.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n").replace("\r", "\\r") + '"'

    translated = string_pattern.sub(repl, script)
    for index, original in enumerate(protected_key_values):
        translated = translated.replace(f"@@SMICUT_MACHINE_KEY_{index}@@", original)
    return translated


def replace_scripts(content: str, mapping: dict[str, str]) -> str:
    pattern = re.compile(r"(<script\b[^>]*>)(?P<script>.*?)(</script>)", re.IGNORECASE | re.DOTALL)

    def repl(match: re.Match[str]) -> str:
        return match.group(1) + replace_js_strings(match.group("script"), mapping) + match.group(3)

    return pattern.sub(repl, content)


def replace_urls(content: str, url_mapping: dict[str, str]) -> str:
    href_pattern = re.compile(r'(?P<prefix>\bhref=)(?P<quote>["\'])(?P<url>.*?)(?P=quote)')

    def repl(match: re.Match[str]) -> str:
        url = html.unescape(match.group("url"))
        target = url_mapping.get(url)
        if not target:
            return match.group(0)
        return f"{match.group('prefix')}{match.group('quote')}{html.escape(target, quote=True)}{match.group('quote')}"

    content = href_pattern.sub(repl, content)
    return content


def replace_material_selector_api_language(content: str, api_language: str) -> str:
    return re.sub(
        r'(?P<prefix>\blanguage:\s*)(["\'])[A-Za-z0-9_-]+(\2)(?P<suffix>,\s*\n\s*history:\s*history)',
        lambda match: f'{match.group("prefix")}"{api_language}"{match.group("suffix")}',
        content,
        count=1,
    )


def enforce_material_category_keys(content: str) -> str:
    material_data_pattern = re.compile(
        r"(?P<prefix>const\s+materialData\s*=\s*\[)(?P<body>.*?)(?P<suffix>\];)",
        flags=re.S,
    )
    match = material_data_pattern.search(content)
    if not match:
        raise ValueError("Could not find materialData in app HTML")

    key_pattern = re.compile(r"(?P<prefix>\bkey\s*:\s*)(?P<quote>['\"])(?P<key>.*?)(?P=quote)")
    found_keys = key_pattern.findall(match.group("body"))
    if len(found_keys) != len(MATERIAL_CATEGORY_KEYS):
        raise ValueError(
            f"Expected {len(MATERIAL_CATEGORY_KEYS)} material keys, found {len(found_keys)}"
        )

    key_index = 0

    def replace_key(key_match: re.Match[str]) -> str:
        nonlocal key_index
        canonical_key = MATERIAL_CATEGORY_KEYS[key_index]
        key_index += 1
        return f"{key_match.group('prefix')}{key_match.group('quote')}{canonical_key}{key_match.group('quote')}"

    canonical_body = key_pattern.sub(replace_key, match.group("body"))
    return content[: match.start()] + match.group("prefix") + canonical_body + match.group("suffix") + content[match.end() :]


def should_lower_material_descriptor(locale: str, descriptor: str) -> bool:
    if not descriptor or lower_first_character(descriptor) == descriptor:
        return False
    forced_prefixes = MATERIAL_DESCRIPTOR_FORCE_LOWER_PREFIXES.get(locale, ())
    if any(descriptor.lstrip().startswith(prefix) for prefix in forced_prefixes):
        return True
    return locale not in MATERIAL_DESCRIPTOR_PRESERVE_CASE_LOCALES


def normalize_material_descriptor(locale: str, group: str, descriptor: str) -> str:
    if group.strip().casefold() == descriptor.strip().casefold():
        return descriptor
    if should_lower_material_descriptor(locale, descriptor):
        return lower_first_character(descriptor)
    return descriptor


def normalize_material_js_descriptor_line(line: str, locale: str) -> str:
    match = re.search(r"group:\s*(['\"])(?P<group>.*?)\1.*?material:\s*(['\"])(?P<material>.*?)\3", line)
    if not match:
        return line
    material = match.group("material")
    normalized = normalize_material_descriptor(locale, match.group("group"), material)
    if normalized == material:
        return line
    return line.replace(f"material: {match.group(3)}{material}{match.group(3)}", f"material: {match.group(3)}{normalized}{match.group(3)}", 1)


def normalize_material_js_descriptors(content: str, locale: str) -> str:
    patterns = [
        re.compile(r"(materials:\s*\[\n)(?P<body>[\s\S]*?)(\n\s*\]\s*\n\s*};)"),
        re.compile(r"(const materialData = \[\n)(?P<body>[\s\S]*?)(\n\s*\];)"),
    ]

    for pattern in patterns:
        def repl(match: re.Match[str]) -> str:
            body = "\n".join(
                normalize_material_js_descriptor_line(line, locale)
                for line in match.group("body").split("\n")
            )
            return match.group(1) + body + match.group(3)

        content = pattern.sub(repl, content, count=1)
    return content


def add_rm_to_material_tensile_values(content: str) -> str:
    content = re.sub(
        r'tensile: "N/mm2 < (?P<value>\d+)"',
        lambda match: f'tensile: "Rm < {match.group("value")} N/mm²"',
        content,
    )
    return re.sub(
        r'tensile: "Rm \(N/mm2\) < (?P<value>\d+)"',
        lambda match: f'tensile: "Rm < {match.group("value")} N/mm²"',
        content,
    )


def normalize_material_category_table(content: str, locale: str) -> str:
    content = re.sub(
        r'(<div class="flex-cell" style="color: #ffffff;">[^<]*(?:<br>))(?!Rm\s)(N/mm²)(</div>)',
        r"\1Rm \2\3",
        content,
    )
    content = content.replace("Tensile Strenght", "Tensile Strength")

    lines = []
    current_group = ""
    group_pattern = re.compile(r'<div class="flex-cell(?: ftext-left)?">(?P<value>[^<]*)</div>')
    descriptor_pattern = re.compile(
        r'(?P<prefix><div class="flex-cell ftext-left" style="flex-basis: 100px; flex-grow: 1;">)'
        r'(?P<value>[^<]*)'
        r'(?P<suffix></div>)'
    )

    for line in content.splitlines():
        group_match = group_pattern.search(line)
        if group_match and "style=" not in line:
            group_value = html.unescape(group_match.group("value")).strip()
            if group_value:
                current_group = group_value

        descriptor_match = descriptor_pattern.search(line)
        if descriptor_match:
            descriptor = html.unescape(descriptor_match.group("value"))
            normalized = normalize_material_descriptor(locale, current_group, descriptor)
            if normalized != descriptor:
                line = (
                    line[: descriptor_match.start()]
                    + descriptor_match.group("prefix")
                    + html.escape(normalized, quote=False)
                    + descriptor_match.group("suffix")
                    + line[descriptor_match.end() :]
                )

        lines.append(line)

    return "\n".join(lines) + ("\n" if content.endswith("\n") else "")


def normalize_material_presentation(content: str, locale: str) -> str:
    content = add_rm_to_material_tensile_values(content)
    content = normalize_material_js_descriptors(content, locale)
    return normalize_material_category_table(content, locale)


def apply_thread_summary_labels(content: str, locale: str) -> str:
    labels = THREAD_SUMMARY_LABELS.get(locale)
    if not labels:
        return content

    replacements = {
        "threadSummaryThread": labels[0],
        "threadSummaryDiameter": labels[1],
        "threadSummaryPitch": labels[2],
        "threadSummaryInternalThread": labels[3],
        "threadSummaryExternalThread": labels[4],
    }
    for key, value in replacements.items():
        content = re.sub(
            rf"({key}:\s*)'[^']*'",
            lambda match: f"{match.group(1)}'{js_escape(value)}'",
            content,
            count=1,
        )
    return content


def apply_thread_profile_labels(content: str, locale: str) -> str:
    labels = THREAD_PROFILE_LABELS.get(locale)
    if not labels:
        return content

    for key, value in labels.items():
        content = re.sub(
            rf"(^\s*{re.escape(key)}:\s*)'[^']*'",
            lambda match: f"{match.group(1)}'{js_escape(value)}'",
            content,
            count=1,
            flags=re.MULTILINE,
        )
        content = re.sub(
            rf'(<option value="{re.escape(key)}">)[^<]*(</option>)',
            lambda match: match.group(1) + html.escape(value, quote=False) + match.group(2),
            content,
        )
    return content


def apply_tap_type_labels(content: str, locale: str) -> str:
    labels = TAP_TYPE_LABELS.get(locale)
    if not labels:
        return content

    js_keys = {
        "tt": "tapTypeTt",
        "tb": "tapTypeTb",
        "tfk": "tapTypeTfk",
    }
    for key, value in labels.items():
        content = re.sub(
            rf"({js_keys[key]}:\s*)'[^']*'",
            lambda match: f"{match.group(1)}'{js_escape(value)}'",
            content,
            count=1,
        )
        content = re.sub(
            rf'(<input[^>]+name="scd-tap-type"[^>]+value="{re.escape(key)}"[^>]*>\s*<span>)[^<]*(</span>)',
            lambda match: match.group(1) + html.escape(value, quote=False) + match.group(2),
            content,
            count=1,
        )
    return content


def replace_start_card_copy(content: str, class_name: str, title: str, description: str) -> str:
    pattern = re.compile(
        rf'(<button type="button" class="scd-start-card scd-start-card-button {re.escape(class_name)}" '
        r'data-scd-section="[^"]+">\s*<strong>)'
        r'[\s\S]*?'
        r'(<span class="scd-card-arrow">➜</span></strong>\s*<span>)'
        r'[\s\S]*?'
        r'(</span>)',
        re.MULTILINE,
    )
    return pattern.sub(
        lambda match: (
            match.group(1)
            + html.escape(title, quote=False)
            + " "
            + match.group(2)
            + html.escape(description, quote=False)
            + match.group(3)
        ),
        content,
        count=1,
    )


def replace_tab_label(content: str, section: str, label: str) -> str:
    pattern = re.compile(
        rf'(<button type="button" class="scd-tab" data-scd-section="{re.escape(section)}">\s*<span>)'
        r"[\s\S]*?"
        r"(</span>\s*</button>)",
        re.MULTILINE,
    )
    return pattern.sub(
        lambda match: match.group(1) + html.escape(label, quote=False) + match.group(2),
        content,
        count=1,
    )


def apply_start_page_labels(content: str, locale: str) -> str:
    labels = START_PAGE_LABELS.get(locale)
    if not labels:
        return content

    content = re.sub(
        r'(<div class="scd-threadmilling-copy">\s*<strong>)[^<]+(</strong>)',
        lambda match: match.group(1) + html.escape(labels["smiprog"], quote=False) + match.group(2),
        content,
        count=1,
    )
    content = replace_tab_label(content, "endmills", labels["milling_title"])
    content = replace_tab_label(content, "drills", labels["drilling_title"])
    content = replace_start_card_copy(
        content,
        "scd-start-card-endmills",
        labels["milling_title"],
        labels["milling_desc"],
    )
    return replace_start_card_copy(
        content,
        "scd-start-card-drills",
        labels["drilling_title"],
        labels["drilling_desc"],
    )


def apply_cjk_typography(content: str, locale: str) -> str:
    if locale not in CJK_UI_LOCALES:
        return content

    root = '<div id="smicut-cutting-data-app">'
    if root not in content:
        raise ValueError("Could not find the SmiCalc root element for CJK typography")

    content = content.replace(
        root,
        '<div id="smicut-cutting-data-app" class="scd-cjk-ui">',
        1,
    )
    return content.rstrip() + "\n\n" + CJK_TYPOGRAPHY_CSS + "\n"


def localize_html(content: str, mapping: dict[str, str], url_mapping: dict[str, str]) -> tuple[str, dict[str, int]]:
    content = replace_urls(content, url_mapping)
    content = replace_br_sequences(content, mapping)
    content = replace_attrs(content, mapping)

    script_blocks: list[str] = []

    def stash_script(match: re.Match[str]) -> str:
        script_blocks.append(match.group(0))
        return f"@@SMICUT_SCRIPT_BLOCK_{len(script_blocks) - 1}@@"

    without_scripts = re.sub(r"<script\b[^>]*>.*?</script>", stash_script, content, flags=re.IGNORECASE | re.DOTALL)
    without_scripts = replace_text_nodes(without_scripts, mapping)

    for index, script in enumerate(script_blocks):
        without_scripts = without_scripts.replace(
            f"@@SMICUT_SCRIPT_BLOCK_{index}@@",
            replace_scripts(script, mapping),
        )

    translated = without_scripts
    stats = {
        "remaining_english_terms": count_remaining_terms(translated, mapping),
    }
    return translated, stats


def count_remaining_terms(content: str, mapping: dict[str, str]) -> int:
    count = 0
    for source, target in mapping.items():
        if source == target or len(source) < 3:
            continue
        if re.search(rf"(?<![A-Za-z0-9_]){re.escape(source)}(?![A-Za-z0-9_])", content):
            count += 1
    return count


def add_material_selector_translations(mapping: dict[str, str], locale: str, source: str) -> None:
    """Localize selector fragments using the main translation workbook."""
    selector_labels = ('Which material will you machine?', 'Enter the material designation or name, and we’ll help you find the right material category.', 'Find material category', 'E.g. AISI 1045 or EN 1.4301', 'The material will be used in all calculators.', 'Material selected', 'Show explanation', 'Hide explanation', 'Ask a follow-up question or enter another material', 'Write a message…', 'Enter your material', 'Open calculator', 'Try again', 'Sending message', 'Send message', 'Material conversation', 'SmiCalc')
    missing = [label for label in selector_labels if label not in mapping]
    if missing:
        raise ValueError(f"Missing Material Selector workbook translations for {locale}: {missing}")
    additions = {label: mapping[label] for label in selector_labels}
    # Preserve spaces around concatenated icon/text fragments.
    for english, translated in list(mapping.items()):
        if english.strip() == english:
            mapping.setdefault(" " + english, " " + translated)
            mapping.setdefault(english + " ", translated + " ")
    selector = source.split('const API_URL =', 1)[1].split('</script>', 1)[0]
    literals = re.findall(r"(['\"])((?:\\.|(?!\1).)*?)\1", selector)
    labels = dict(additions)
    for label in ("Recommended material category", "Use in SmiCalc", "Material selected", "Selected material category: ", "SmiCalc"):
        if label in mapping:
            labels[label] = mapping[label]
    labels["Material selected: "] = additions["Material selected"] + ": "
    for quote, literal in literals:
        # Embedded markup needs HTML escaping as well as JS escaping (done later).
        if '<' not in literal:
            continue
        translated = literal
        for english, target in sorted(labels.items(), key=lambda item: -len(item[0])):
            translated = translated.replace(english, html.escape(target, quote=False))
        if translated != literal:
            mapping[literal] = translated
    mapping["Material selected: "] = additions["Material selected"] + ": "


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("locale", choices=sorted(LOCALES))
    parser.add_argument(
        "--app-only",
        action="store_true",
        help="Regenerate only the SmiCalc cutting-data app HTML and leave material-category files untouched.",
    )
    args = parser.parse_args()

    config = LOCALES[args.locale]
    mapping = load_translation_map(config["target_column"])
    mapping.update(UI_TRANSLATION_OVERRIDES.get(args.locale, {}))
    url_mapping = load_url_map(config["url_column"])
    url_mapping.update(config.get("url_overrides", {}))
    source_dir = ROOT / "Translation" / config["source_folder"]
    target_dir = ROOT / "Translation" / config["folder"]
    target_dir.mkdir(parents=True, exist_ok=True)

    outputs = []
    source_files = SOURCE_FILES.items()
    if args.app_only:
        source_files = [
            ("smicut cutting data app-eng.html", SOURCE_FILES["smicut cutting data app-eng.html"])
        ]

    for source_name, target_template in source_files:
        source_path = source_dir / source_name
        target_path = target_dir / target_template.format(suffix=config["suffix"])
        content = source_path.read_text(encoding="utf-8")
        if source_name == "smicut cutting data app-eng.html":
            add_material_selector_translations(mapping, args.locale, content)
        localized, stats = localize_html(content, mapping, url_mapping)
        localized = replace_material_selector_api_language(localized, config["api_language"])
        if source_name == "smicut cutting data app-eng.html":
            localized = enforce_material_category_keys(localized)
            localized = apply_start_page_labels(localized, args.locale)
            localized = apply_thread_summary_labels(localized, args.locale)
            localized = apply_thread_profile_labels(localized, args.locale)
            localized = apply_tap_type_labels(localized, args.locale)
            localized = normalize_material_presentation(localized, args.locale)
            localized = apply_cjk_typography(localized, args.locale)
        target_path.write_text(localized, encoding="utf-8")
        outputs.append(
            {
                "source": str(source_path.relative_to(ROOT)),
                "target": str(target_path.relative_to(ROOT)),
                "remaining_english_term_count": stats["remaining_english_terms"],
            }
        )

    print(json.dumps({"locale": args.locale, "outputs": outputs}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
